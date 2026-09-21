"""Many-folder ZIP previews stay bounded and preserve choices across pages."""
import hashlib
import io
import re
import zipfile
from pathlib import Path

from sqlalchemy import event
from tests.test_zip_import_review import app, client, archive
from tests.test_importer import _upload, _commit, _job


def populate(app,count=251):
    from app.extensions import db
    from app.models import Matter
    with app.app_context():
        db.session.add_all([Matter(number=f'F-{i:04d}',name=f'Synthetic folder matter {i}',client_id=1)
                            for i in range(count)])
        db.session.commit()


def field(folder):return 'folder_number_'+hashlib.md5(folder.encode()).hexdigest()


def test_many_folders_do_not_repeat_every_matter_list(app,client):
    populate(app)
    raw=archive([(f'F-{i:04d}/doc.txt',b'Synthetic document') for i in range(101)])
    token=_upload(client,'documents',raw,'folders.zip')
    html=client.get('/import/preview/'+token).get_data(as_text=True)
    assert html.count('<option')<=260
    assert len(re.findall(r'name="folder_(?:number_)?[a-f0-9]+"',html))<=25
    assert 'Page 1 of 5' in html and 'Import all 101 files' in html


def test_folder_matching_reads_matters_once(app,client):
    populate(app,60)
    from app.extensions import db
    selects=[]
    def record(conn,cursor,statement,*args):
        if statement.lstrip().upper().startswith('SELECT') and 'FROM matters' in statement:
            selects.append(statement)
    with app.app_context():
        engine=db.engine;event.listen(engine,'before_cursor_execute',record)
    try:
        _upload(client,'documents',archive([(f'F-{i:04d}/d.txt',b'Synthetic') for i in range(51)]),'many.zip')
    finally:event.remove(engine,'before_cursor_execute',record)
    assert len(selects)<=2,len(selects)


def test_save_navigation_keeps_other_page_choices_and_imports_all_folders(app,client):
    from app.models import Document,Matter
    populate(app,60)
    token=_upload(client,'documents',archive([(f'F-{i:04d}/file-{i}.txt',b'Synthetic') for i in range(51)]),'pages.zip')
    path='/import/preview/'+token
    r=client.post(path,data={'_csrf':client.tok,'do':'page:2',field('F-0000'):''})
    assert r.status_code==302 and 'page=2' in r.location
    r=client.post(path+'?page=2',data={'_csrf':client.tok,'do':'page:1',field('F-0025'):'F-0059'})
    assert r.status_code==302
    first=client.get(r.location).get_data(as_text=True)
    assert re.search('name="'+field('F-0000')+'"[^>]*value=""',first)
    jid=_commit(client,token)
    j=_job(app,jid);assert (j['created'],j['skipped'],len(j['errors']))==(50,1,0),j
    with app.app_context():
        assert Document.query.filter_by(name='file-0.txt').count()==0
        d=Document.query.filter_by(name='file-25.txt').one()
        assert d.matter_id==Matter.query.filter_by(number='F-0059').one().id
        assert Document.query.filter_by(name='file-50.txt').count()==1


def test_unknown_typed_matter_cannot_silently_use_auto_match(app,client):
    from app.models import ImportJob,Document
    token=_upload(client,'documents',archive([('ZIP-1/doc.txt',b'Synthetic')]),'unknown.zip')
    response=client.post('/import/preview/'+token,data={'_csrf':client.tok,'do':'commit',field('ZIP-1'):'does-not-exist'})
    assert response.status_code==200 and b'Unknown matter number' in response.data
    with app.app_context():assert ImportJob.query.count()==0 and Document.query.count()==0


def test_blank_matter_name_does_not_match_unrelated_folder(app,client):
    from app.extensions import db
    from app.models import Matter
    with app.app_context():
        db.session.add(Matter(number='BLANK-1',name='',client_id=1));db.session.commit()
    token=_upload(client,'documents',archive([('Entirely unrelated folder/doc.txt',b'Synthetic')]),'unrelated.zip')
    job=_job(app,_commit(client,token))
    assert job['created']==0 and len(job['errors'])==1,job


def test_mixed_pdf_word_and_text_import_preserves_bytes_and_search_text(app,client):
    from docx import Document as WordDocument
    from fpdf import FPDF
    from app.models import Document
    pdf=FPDF()
    for i in range(12):
        pdf.add_page();pdf.set_font('Helvetica',size=12);pdf.multi_cell(0,8,f'Synthetic PDF evidence page {i+1}. Amount 1234.56.')
    word=WordDocument();word.add_heading('Synthetic Word evidence',0)
    word.add_paragraph('Witness corrected the time from 09:15 to 10:15. Amount 1234.56.')
    word_bytes=io.BytesIO();word.save(word_bytes)
    contents={'evidence.pdf':bytes(pdf.output()),'evidence.docx':word_bytes.getvalue(),
              'notes.txt':b'Synthetic plain text evidence. Do not send.',
              'message.eml':b'From: synthetic@example.test\r\nSubject: Synthetic email\r\n\r\nSynthetic email body.'}
    raw=archive([(f'ZIP-1/{name}',body) for name,body in contents.items()])
    j=_job(app,_commit(client,_upload(client,'documents',raw,'mixed.zip')))
    assert j['created']==4 and not j['errors'],j
    with app.app_context():
        for d in Document.query.all():
            assert (Path(app.config['UPLOAD_DIR'])/d.path).read_bytes()==contents[d.name]
            assert 'Synthetic' in d.extracted_text
            if d.name=='evidence.pdf':assert 'page 12' in d.extracted_text
            if d.name=='evidence.docx':assert '09:15 to 10:15' in d.extracted_text
            assert not d.shared_to_portal
    j=_job(app,_commit(client,_upload(client,'documents',raw,'mixed.zip')))
    assert j['created']==0 and j['skipped']==4 and not j['errors']


def test_duplicate_matter_names_require_an_explicit_choice(app,client):
    from app.extensions import db
    from app.models import Matter
    with app.app_context():
        db.session.add_all([Matter(number='SAME-1',name='Shared matter name',client_id=1),
                            Matter(number='SAME-2',name='Shared matter name',client_id=1)])
        db.session.commit()
    token=_upload(client,'documents',archive([('Shared matter name/doc.txt',b'Synthetic')]),'ambiguous-matter.zip')
    job=_job(app,_commit(client,token))
    assert job['created']==0 and len(job['errors'])==1,job
